"""Readers for the CEN spreadsheets.

The generation workbook is ~100 MB, so it is read in streaming mode and
aggregated on the fly to (hour, region, type, subtype). Nothing is kept per
unit, which keeps memory use small.
"""

from __future__ import annotations

import datetime as dt
import logging
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

log = logging.getLogger(__name__)

# Canonical region names used across the package.
REGIONS = {
    "arica y parinacota": "Arica y Parinacota",
    "tarapaca": "Tarapacá",
    "antofagasta": "Antofagasta",
    "atacama": "Atacama",
    "coquimbo": "Coquimbo",
    "valparaiso": "Valparaíso",
    "metropolitana": "Metropolitana",
    "ohiggins": "O'Higgins",
    "libertador general bernardo ohiggins": "O'Higgins",
    "maule": "Maule",
    "nuble": "Ñuble",
    "biobio": "Biobío",
    "la araucania": "La Araucanía",
    "araucania": "La Araucanía",
    "los rios": "Los Ríos",
    "los lagos": "Los Lagos",
    "aysen": "Aysén",
    "magallanes": "Magallanes",
}


def _key(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    text = text.lower().strip()
    text = re.sub(r"^region\s+(de\s+la\s+|del\s+|de\s+)?", "", text)
    text = text.replace("-", "").replace("'", "").replace("’", "")
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+y\s+la\s+antartica\s+chilena$", "", text)
    return text


def normalize_region(name: str) -> str:
    """Map the many spellings found in CEN files to one canonical region name."""
    key = _key(name)
    if key in REGIONS:
        return REGIONS[key]
    raise ValueError(f"unknown region name: {name!r}")


class _BadValues:
    """Counts cells that could not be read as numbers."""

    def __init__(self) -> None:
        self.count = 0
        self.examples: list[str] = []


def _to_float(value, bad: _BadValues | None = None) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text in {"-", "(en blanco)"}:
        return 0.0
    try:
        return float(text.replace(",", "."))
    except ValueError:
        if bad is not None:
            bad.count += 1
            if len(bad.examples) < 5:
                bad.examples.append(text)
        return 0.0


MONTHS = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
          "jul": 7, "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12}


def _to_date(value, year: int | None = None) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):  # Excel serial number
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(value))).date()
    if isinstance(value, str) and value.strip():
        text = value.strip().lower()
        m = re.fullmatch(r"(\d{1,2})-([a-z]{3,4})\.?", text)
        if m:  # "01-ene": the year is not in the cell
            if year is None:
                raise ValueError(f"date {value!r} has no year; pass year=")
            return dt.date(year, MONTHS[m.group(2)[:3]], int(m.group(1)))
        if re.match(r"\d{4}-\d{1,2}-\d{1,2}", text):
            return pd.to_datetime(text).date()
        return pd.to_datetime(text, dayfirst=True).date()
    return None


# Column names used by the two layouts CEN has published so far.
LAYOUTS = [
    # 2025 edition onwards
    {"region": "Región", "tipo": "Tipo", "subtipo": "Subtipo", "date": "Fecha", "hour": "Hora {h}"},
    # 2024 edition
    {"region": "region", "tipo": "tipo_fuente", "subtipo": "subtipo_fuente",
     "date": "_ops_por_hora_new.fecha", "hour": "{h}"},
]


def _match_layout(cells: list[str]) -> dict | None:
    for layout in LAYOUTS:
        needed = [layout["region"], layout["tipo"], layout["subtipo"], layout["date"],
                  layout["hour"].format(h=1), layout["hour"].format(h=24)]
        if all(n in cells for n in needed):
            return layout
    return None


def read_generation_by_unit(path: str | Path, year: int | None = None) -> pd.DataFrame:
    """Read a CEN "hist_gen_de_energia_por_central" workbook.

    Returns a long table with columns ``time, region, tipo, subtipo, mwh``.
    ``time`` is the start of the hour in local clock time: hour 1 of a day
    is 00:00-01:00. Days are indexed with 24 hours. Some editions add a 25th
    hour column for the day daylight saving time ends; that energy cannot be
    placed on a 24-hour clock and is reported in ``df.attrs["hour25_mwh"]``
    instead. ``year`` is needed for editions whose dates have no year.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        layout = None
        for row in rows:
            cells = [str(c).strip() if c is not None else "" for c in row]
            layout = _match_layout(cells)
            if layout:
                header = cells
                break
        if layout is None:
            raise ValueError(f"{path}: no known header layout found")
        i_region = header.index(layout["region"])
        i_tipo = header.index(layout["tipo"])
        i_sub = header.index(layout["subtipo"])
        i_date = header.index(layout["date"])
        i_hours = [header.index(layout["hour"].format(h=h)) for h in range(1, 25)]
        h25 = layout["hour"].format(h=25)
        i_h25 = header.index(h25) if h25 in header else None

        acc: dict[tuple, list[float]] = defaultdict(lambda: [0.0] * 24)
        unknown_regions: set[str] = set()
        hour25 = 0.0
        bad = _BadValues()
        for row in rows:
            if row is None or len(row) <= i_hours[-1]:
                continue
            day = _to_date(row[i_date], year)
            if day is None:
                continue
            raw_region = row[i_region]
            try:
                region = normalize_region(raw_region) if raw_region else "Unknown"
            except ValueError:
                unknown_regions.add(str(raw_region))
                region = "Unknown"
            key = (day, region, str(row[i_tipo]).strip(), str(row[i_sub]).strip())
            values = acc[key]
            for h, idx in enumerate(i_hours):
                values[h] += _to_float(row[idx], bad)
            if i_h25 is not None and i_h25 < len(row):
                hour25 += _to_float(row[i_h25], bad)
    finally:
        wb.close()

    records = []
    for (day, region, tipo, sub), values in acc.items():
        base = dt.datetime.combine(day, dt.time())
        for h, v in enumerate(values):
            records.append((base + dt.timedelta(hours=h), region, tipo, sub, v))
    df = pd.DataFrame(records, columns=["time", "region", "tipo", "subtipo", "mwh"])
    df = df.sort_values(["time", "region", "tipo", "subtipo"]).reset_index(drop=True)
    df.attrs["unknown_regions"] = sorted(unknown_regions)
    df.attrs["hour25_mwh"] = hour25
    df.attrs["non_numeric_cells"] = bad.count
    if bad.count:
        log.warning("%s: %d non-numeric hour values read as 0, e.g. %s", path, bad.count, bad.examples)
    return df


def read_capacity_by_region(path: str | Path) -> pd.DataFrame:
    """Read the CEN "hist_cap_inst_por_region_y_tecno" workbook.

    The sheet stacks one block per year, two blocks side by side. Each block
    starts with a header row ``<year>, Geotérmica, Termosolar, ...`` followed
    by one row per region and a "Total" row.

    Returns a long table with columns ``year, region, technology, mw``.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = next(s for s in wb.worksheets if s.title.lower().startswith("capacidad"))
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    records = []
    for ri, row in enumerate(grid):
        for ci, cell in enumerate(row):
            if not (isinstance(cell, int) and 1990 < cell < 2100):
                continue
            if ci + 1 >= len(row) or row[ci + 1] != "Geotérmica":
                continue
            techs = []
            for c in row[ci + 1 :]:
                if c is None or isinstance(c, int):
                    break
                techs.append(str(c).strip())
            for r in grid[ri + 1 :]:
                name = r[ci] if ci < len(r) else None
                if name is None or str(name).strip().lower() == "total":
                    break
                region = normalize_region(name)
                for k, tech in enumerate(techs):
                    if tech == "Total":
                        continue
                    records.append((cell, region, tech, _to_float(r[ci + 1 + k])))
    if not records:
        raise ValueError(f"{path}: no capacity blocks found")
    return pd.DataFrame(records, columns=["year", "region", "technology", "mw"])
