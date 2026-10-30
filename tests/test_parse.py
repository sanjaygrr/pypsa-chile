import datetime as dt

import pytest

from conftest import FIXTURES
from pypsa_chile.parse import normalize_region, read_capacity_by_region, read_generation_by_unit


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("Región del Bío-Bío", "Biobío"),
        ("Biobío", "Biobío"),
        ("Región Metropolitana", "Metropolitana"),
        ("Región de O'Higgins", "O'Higgins"),
        ("los Lagos", "Los Lagos"),
        ("Región de la Araucanía", "La Araucanía"),
        ("Arica Y Parinacota", "Arica y Parinacota"),
        ("Región de Ñuble", "Ñuble"),
    ],
)
def test_normalize_region(raw, expected):
    assert normalize_region(raw) == expected


def test_normalize_region_unknown():
    with pytest.raises(ValueError):
        normalize_region("Atlantis")


def test_read_generation_by_unit():
    df = read_generation_by_unit(FIXTURES / "generation_by_unit_excerpt.xlsx")
    assert list(df.columns) == ["time", "region", "tipo", "subtipo", "mwh"]
    assert df["time"].min() == dt.datetime(2025, 10, 6, 0)
    assert df["time"].max() == dt.datetime(2025, 10, 6, 23)
    # 24 hours for every (region, tipo, subtipo) group
    assert (df.groupby(["region", "tipo", "subtipo"]).size() == 24).all()
    assert "Unknown" not in set(df["region"])
    assert df.attrs["unknown_regions"] == []
    # spellings are merged: no "Región ..." variants survive
    assert not df["region"].str.startswith("Región").any()


def test_read_generation_hours_match_row_totals():
    """Sum over the long table equals the sum of the 24 hour columns in the sheet."""
    from openpyxl import load_workbook

    wb = load_workbook(FIXTURES / "generation_by_unit_excerpt.xlsx", read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    header = list(rows[0])
    i1 = header.index("Hora 1")
    expected = sum(float(v or 0) for r in rows[1:] for v in r[i1 : i1 + 24])
    df = read_generation_by_unit(FIXTURES / "generation_by_unit_excerpt.xlsx")
    assert df["mwh"].sum() == pytest.approx(expected)


def test_read_capacity_by_region(capacity):
    years = sorted(capacity["year"].unique())
    assert years[0] == 2015 and 2024 in years
    y = capacity[capacity["year"] == 2024]
    assert set(y["technology"]) >= {"Solar", "Eólica", "Carbón", "Gas Natural", "Hidroeléctrica"}
    assert y["region"].nunique() == 14  # SEN regions, Arica to Los Lagos
    # The sheet's own "Total" row for 2024 is 36,777.9 MW.
    assert y["mw"].sum() == pytest.approx(36777.92, abs=0.1)


def test_read_capacity_bad_file(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Capacidad vacia"
    wb.save(tmp_path / "empty.xlsx")
    with pytest.raises(ValueError):
        read_capacity_by_region(tmp_path / "empty.xlsx")
