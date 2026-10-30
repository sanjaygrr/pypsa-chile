"""Rebuild the test fixtures from the real CEN files.

Run from the repository root after `pypsa-chile fetch` with the default
config:

    python tests/fixtures/make_fixtures.py

The fixtures are small excerpts of public CEN data; nothing is synthetic.
"""

from __future__ import annotations

import datetime as dt
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

HERE = Path(__file__).parent
RAW = Path("data/raw")
PROCESSED = Path("data/processed")

GEN_XLSX = RAW / "CEN-hist_gen_de_energia_por_central_2025.xlsx"
GEN_XLSX_2024 = RAW / "CEN-hist_gen_de_energia_por_central_2024.xlsx"
CAP_XLSX = RAW / "CEN-hist_cap_inst_por_region_y_tecno.xlsx"


def excerpt_generation_xlsx(day: dt.date, per_label: int = 2) -> None:
    """Keep the header and a few units of each (region spelling, type) on one day."""
    src = load_workbook(GEN_XLSX, read_only=True, data_only=True)
    out = Workbook()
    ws_out = out.active
    ws_out.title = "Principal"
    seen: dict[tuple, int] = {}
    rows = src.worksheets[0].iter_rows(values_only=True)
    ws_out.append(list(next(rows)))
    for row in rows:
        d = row[7]
        if not isinstance(d, dt.datetime) or d.date() != day:
            continue
        key = (row[2], row[5], row[6])
        if seen.get(key, 0) >= per_label:
            continue
        seen[key] = seen.get(key, 0) + 1
        ws_out.append(list(row))
    src.close()
    out.save(HERE / "generation_by_unit_excerpt.xlsx")


def excerpt_generation_xlsx_2024(n_rows: int = 400) -> None:
    """First rows of the 2024 edition, which uses a different layout."""
    src = load_workbook(GEN_XLSX_2024, read_only=True, data_only=True)
    out = Workbook()
    ws_out = out.active
    ws_out.title = "Hoja2"
    for i, row in enumerate(src.worksheets[0].iter_rows(values_only=True)):
        if i > n_rows:
            break
        ws_out.append(list(row))
    src.close()
    out.save(HERE / "generation_by_unit_2024_excerpt.xlsx")


def excerpt_processed(start: str, end: str) -> None:
    gen = pd.read_csv(PROCESSED / "generation_2025.csv.gz", parse_dates=["time"])
    sel = gen[(gen["time"] >= start) & (gen["time"] < end)]
    sel.to_csv(HERE / "generation_processed_excerpt.csv.gz", index=False, compression="gzip")


if __name__ == "__main__":
    shutil.copy(CAP_XLSX, HERE / "capacity_by_region.xlsx")
    excerpt_processed("2025-10-03", "2025-10-07")
    excerpt_generation_xlsx(dt.date(2025, 10, 6))
    excerpt_generation_xlsx_2024()
